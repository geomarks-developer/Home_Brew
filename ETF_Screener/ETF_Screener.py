from pathlib import Path
import time
import warnings

import numpy as np
import pandas as pd
import yfinance as yf


# ============================================================
# USER SETTINGS
# ============================================================

TICKERS = [
    "SPY", "VOO", "IVV", "VTI", "QQQ", "IWM",
    "SCHD", "VIG", "VYM", "DGRO", "VTV", "VUG",
    "XLK", "XLF", "XLE", "XLV", "XLI", "XLY", "XLP", "XLU", "XLRE", "XLB",
    "VEA", "VWO", "VXUS", "EFA", "EEM",
    "BND", "AGG", "SGOV", "TLT", "IEF", "HYG", "LQD",
    "VNQ", "GLD", "SLV", "DBC",
]


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "ETF_Screener.xlsx"
OUTPUT_SHEET = "ETF Screener"
PRICE_PERIOD = "5y"
RISK_FREE_RATE = 0.04
METADATA_DELAY_SECONDS = 0.25


def safe_number(value):
    try:
        number = float(value)
        return number if np.isfinite(number) else np.nan
    except (TypeError, ValueError):
        return np.nan


def normalize_decimal(value):
    number = safe_number(value)
    if pd.isna(number):
        return np.nan
    return number / 100 if abs(number) > 1 else number


def first_available(dictionary, keys, default=np.nan):
    for key in keys:
        value = dictionary.get(key)
        if value is not None and value != "":
            return value
    return default


def extract_close_prices(downloaded, tickers):
    if downloaded.empty:
        return pd.DataFrame()

    prices = None
    if isinstance(downloaded.columns, pd.MultiIndex):
        level_zero = downloaded.columns.get_level_values(0)
        level_one = downloaded.columns.get_level_values(1)
        if "Close" in level_zero:
            prices = downloaded["Close"].copy()
        elif "Adj Close" in level_zero:
            prices = downloaded["Adj Close"].copy()
        elif "Close" in level_one:
            prices = downloaded.xs("Close", axis=1, level=1).copy()
        elif "Adj Close" in level_one:
            prices = downloaded.xs("Adj Close", axis=1, level=1).copy()
    else:
        field = "Close" if "Close" in downloaded.columns else "Adj Close"
        if field in downloaded.columns:
            prices = downloaded[[field]].copy()
            if len(tickers) == 1:
                prices.columns = [tickers[0]]

    if prices is None:
        raise ValueError("Could not locate Close prices in yfinance output.")
    if isinstance(prices, pd.Series):
        prices = prices.to_frame(name=tickers[0])
    prices.columns = [str(column).upper() for column in prices.columns]
    return prices.sort_index()


def trailing_return(prices, trading_days):
    clean = prices.dropna()
    if len(clean) <= trading_days:
        return np.nan
    start_price = clean.iloc[-trading_days - 1]
    return clean.iloc[-1] / start_price - 1 if start_price > 0 else np.nan


def calendar_ytd_return(prices):
    clean = prices.dropna()
    if len(clean) < 2:
        return np.nan
    current_year = clean.index[-1].year
    prior_year = clean[clean.index.year < current_year]
    current_year_data = clean[clean.index.year == current_year]
    if current_year_data.empty:
        return np.nan
    start_price = prior_year.iloc[-1] if not prior_year.empty else current_year_data.iloc[0]
    return current_year_data.iloc[-1] / start_price - 1 if start_price > 0 else np.nan


def annualized_return(prices):
    clean = prices.dropna()
    if len(clean) < 2 or clean.iloc[0] <= 0:
        return np.nan
    years = (clean.index[-1] - clean.index[0]).days / 365.25
    return (clean.iloc[-1] / clean.iloc[0]) ** (1 / years) - 1 if years > 0 else np.nan


def annualized_volatility(prices):
    returns = prices.dropna().pct_change().dropna()
    return returns.std() * np.sqrt(252) if len(returns) >= 2 else np.nan


def maximum_drawdown(prices):
    clean = prices.dropna()
    if clean.empty:
        return np.nan
    return (clean / clean.cummax() - 1).min()


def sharpe_ratio(prices, risk_free_rate=0.04):
    returns = prices.dropna().pct_change().dropna()
    if len(returns) < 2 or returns.std() == 0:
        return np.nan
    daily_rf = (1 + risk_free_rate) ** (1 / 252) - 1
    return (returns.mean() - daily_rf) / returns.std() * np.sqrt(252)


def downside_deviation(prices):
    negative = prices.dropna().pct_change().dropna()
    negative = negative[negative < 0]
    return negative.std() * np.sqrt(252) if len(negative) >= 2 else np.nan


def distance_from_high(prices, lookback_days=252):
    clean = prices.dropna().iloc[-lookback_days:]
    if clean.empty or clean.max() <= 0:
        return np.nan
    return clean.iloc[-1] / clean.max() - 1


def moving_average_status(prices, window):
    clean = prices.dropna()
    if len(clean) < window:
        return np.nan
    return int(clean.iloc[-1] > clean.rolling(window).mean().iloc[-1])


def get_fund_metadata(symbol):
    ticker = yf.Ticker(symbol)
    try:
        info = ticker.info or {}
    except Exception as error:
        print(f"  Metadata warning for {symbol}: {error}")
        info = {}
    try:
        fast_info = dict(ticker.fast_info)
    except Exception:
        fast_info = {}

    return {
        "Name": first_available(info, ["longName", "shortName"], symbol),
        "Category": first_available(info, ["category"], ""),
        "Fund Family": info.get("fundFamily", ""),
        "Exchange": first_available(info, ["exchange", "fullExchangeName"], ""),
        "Currency": info.get("currency", ""),
        "Quote Type": info.get("quoteType", ""),
        "AUM": safe_number(first_available(info, ["totalAssets", "netAssets"])),
        "Expense Ratio": normalize_decimal(first_available(info, ["annualReportExpenseRatio", "netExpenseRatio"])),
        "Dividend Yield": normalize_decimal(first_available(info, ["yield", "dividendYield", "trailingAnnualDividendYield"])),
        "Three-Year Beta": safe_number(info.get("beta3Year")),
        "Morningstar Rating": safe_number(info.get("morningStarOverallRating")),
        "Average Volume": safe_number(first_available(info, ["averageVolume", "averageDailyVolume10Day"])),
        "NAV": safe_number(first_available(info, ["navPrice", "nav"])),
        "Previous Close": safe_number(first_available(fast_info, ["previous_close", "previousClose"])),
    }


def download_price_history(tickers):
    print("\nDownloading ETF price history...")
    downloaded = yf.download(
        tickers=tickers,
        period=PRICE_PERIOD,
        interval="1d",
        auto_adjust=True,
        actions=False,
        group_by="column",
        threads=True,
        progress=False,
    )
    prices = extract_close_prices(downloaded, tickers)
    if prices.empty:
        raise RuntimeError("No price data was returned by yfinance.")
    return prices


def build_screener(tickers):
    tickers = list(dict.fromkeys(s.strip().upper() for s in tickers if s.strip()))
    prices = download_price_history(tickers)
    records = []
    print("\nRetrieving ETF metadata and calculating metrics...")

    for position, symbol in enumerate(tickers, start=1):
        print(f"[{position}/{len(tickers)}] {symbol}")
        if symbol not in prices.columns:
            print(f"  Skipped: no price column returned for {symbol}")
            continue
        series = prices[symbol].dropna()
        if len(series) < 30:
            print(f"  Skipped: insufficient price history for {symbol}")
            continue

        metadata = get_fund_metadata(symbol)
        last_price = series.iloc[-1]
        previous_price = series.iloc[-2] if len(series) >= 2 else np.nan
        records.append({
            "Ticker": symbol,
            **metadata,
            "Last Price": last_price,
            "Price Date": series.index[-1].date(),
            "Daily Return": last_price / previous_price - 1 if pd.notna(previous_price) and previous_price > 0 else np.nan,
            "1M Return": trailing_return(series, 21),
            "3M Return": trailing_return(series, 63),
            "6M Return": trailing_return(series, 126),
            "YTD Return": calendar_ytd_return(series),
            "1Y Return": trailing_return(series, 252),
            "3Y Return": trailing_return(series, 756),
            "5Y Annualized Return": annualized_return(series),
            "Annualized Volatility": annualized_volatility(series),
            "Downside Deviation": downside_deviation(series),
            "Maximum Drawdown": maximum_drawdown(series),
            "Sharpe Ratio": sharpe_ratio(series, RISK_FREE_RATE),
            "Distance From 52W High": distance_from_high(series),
            "Above 50-Day MA": moving_average_status(series, 50),
            "Above 200-Day MA": moving_average_status(series, 200),
            "History Observations": len(series),
        })
        if METADATA_DELAY_SECONDS > 0:
            time.sleep(METADATA_DELAY_SECONDS)

    if not records:
        raise RuntimeError("No ETF records were created. Check the ticker list and internet connection.")
    return pd.DataFrame(records)


def percentile_score(series, higher_is_better=True):
    numeric = pd.to_numeric(series, errors="coerce")
    if numeric.notna().sum() == 0:
        return pd.Series(np.nan, index=series.index, dtype=float)
    return numeric.rank(pct=True, ascending=higher_is_better) * 100


def add_rankings(df):
    result = df.copy()
    score_specs = {
        "1M Score": ("1M Return", True),
        "3M Score": ("3M Return", True),
        "6M Score": ("6M Return", True),
        "1Y Score": ("1Y Return", True),
        "Sharpe Score": ("Sharpe Ratio", True),
        "Volatility Score": ("Annualized Volatility", False),
        "Drawdown Score": ("Maximum Drawdown", True),
        "Expense Score": ("Expense Ratio", False),
        "Yield Score": ("Dividend Yield", True),
    }
    for score_name, (source, higher_is_better) in score_specs.items():
        result[score_name] = percentile_score(result[source], higher_is_better)

    result["Momentum Score"] = result[["1M Score", "3M Score", "6M Score", "1Y Score"]].mean(axis=1, skipna=True)
    result["Risk Score"] = result[["Sharpe Score", "Volatility Score", "Drawdown Score"]].mean(axis=1, skipna=True)
    result["Trend Score"] = pd.DataFrame({
        "MA50": result["Above 50-Day MA"] * 100,
        "MA200": result["Above 200-Day MA"] * 100,
    }).mean(axis=1, skipna=True)

    result["Overall Score"] = (
        result["Momentum Score"] * 0.45
        + result["Risk Score"] * 0.25
        + result["Trend Score"] * 0.15
        + result["Expense Score"] * 0.10
        + result["Yield Score"] * 0.05
    )
    result["Overall Rank"] = result["Overall Score"].rank(method="min", ascending=False).astype("Int64")
    result["Signal"] = np.select(
        [
            (result["Overall Score"] >= 75) & (result["Above 200-Day MA"] == 1),
            (result["Overall Score"] >= 55) & (result["Above 200-Day MA"] == 1),
            result["Above 200-Day MA"] == 0,
        ],
        ["Strong", "Positive", "Weak Trend"],
        default="Neutral",
    )
    return result.sort_values(["Overall Rank", "Ticker"]).reset_index(drop=True)


def format_excel_sheet(file_path, sheet_name, df):
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(file_path)
    worksheet = workbook[sheet_name]
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    percentage_columns = {
        "Expense Ratio", "Dividend Yield", "Daily Return", "1M Return", "3M Return",
        "6M Return", "YTD Return", "1Y Return", "3Y Return", "5Y Annualized Return",
        "Annualized Volatility", "Downside Deviation", "Maximum Drawdown", "Distance From 52W High",
    }
    money_columns = {"Last Price", "NAV", "Previous Close"}
    score_columns = {column for column in df.columns if column.endswith("Score")}

    for column_index, column_name in enumerate(df.columns, start=1):
        column_letter = get_column_letter(column_index)
        if column_name in percentage_columns:
            number_format = "0.00%"
        elif column_name in money_columns:
            number_format = "$#,##0.00"
        elif column_name == "AUM":
            number_format = '$#,##0,," M"'
        elif column_name in {"Average Volume", "History Observations", "Overall Rank"}:
            number_format = "#,##0"
        elif column_name in score_columns:
            number_format = "0.0"
        elif column_name in {"Sharpe Ratio", "Three-Year Beta", "Morningstar Rating"}:
            number_format = "0.00"
        else:
            number_format = None

        if number_format:
            for cell in worksheet[column_letter][1:]:
                cell.number_format = number_format

        max_length = max(
            [len(str(column_name))]
            + [len(str(cell.value)) for cell in worksheet[column_letter] if cell.value is not None]
        )
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 11), 35)

    if "Overall Score" in df.columns:
        score_letter = get_column_letter(df.columns.get_loc("Overall Score") + 1)
        worksheet.conditional_formatting.add(
            f"{score_letter}2:{score_letter}{worksheet.max_row}",
            ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            ),
        )
    workbook.save(file_path)


def export_to_excel(df):
    #OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    options = {"engine": "openpyxl", "mode": "a" if OUTPUT_FILE.exists() else "w"}
    if OUTPUT_FILE.exists():
        options["if_sheet_exists"] = "replace"
    with pd.ExcelWriter(OUTPUT_FILE, **options) as writer:
        df.to_excel(writer, sheet_name=OUTPUT_SHEET, index=False)
    format_excel_sheet(OUTPUT_FILE, OUTPUT_SHEET, df)
    print(f"\nExcel file saved to:\n{OUTPUT_FILE}")


def main():
    warnings.filterwarnings("ignore", category=FutureWarning)
    print("=" * 60)
    print("YFINANCE ETF SCREENER")
    print("=" * 60)
    screener = add_rankings(build_screener(TICKERS))
    export_to_excel(screener)

    display_columns = [
        "Overall Rank", "Ticker", "Name", "Last Price", "1M Return", "3M Return",
        "6M Return", "1Y Return", "Annualized Volatility", "Sharpe Ratio",
        "Expense Ratio", "Dividend Yield", "Overall Score", "Signal",
    ]
    print("\nTop-ranked ETFs:\n")
    print(screener[[c for c in display_columns if c in screener.columns]].head(15).to_string(index=False))


if __name__ == "__main__":
    main()
