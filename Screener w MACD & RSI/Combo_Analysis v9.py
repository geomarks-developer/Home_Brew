from __future__ import annotations

import math
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
import plotly.express as px
import yfinance as yf
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORKBOOK_NAME = "Combo_Analysis.xlsx"
TICKERS_SHEET = "Tickers"
HTML_NAME = "investment_dashboard.html"
MONTH_LOOKBACK = 21
MAX_PEG, MAX_FORWARD_PE = 2.5, 30.0
MIN_ROE, MIN_MARGIN = 0.05, 0.05
MAX_DEBT_EQUITY, MIN_UPSIDE = 150.0, 0.0
MIN_DIVIDEND_YIELD, MAX_PAYOUT = 0.02, 0.70
MAX_EXPENSE_RATIO, MIN_AUM = 0.005, 500_000_000
MIN_RSI, MAX_RSI = 40.0, 70.0
PAUSE = 0.05

SECTOR_ETFS = {
    "Communication Services": "XLC", "Consumer Discretionary": "XLY",
    "Consumer Staples": "XLP", "Energy": "XLE", "Financials": "XLF",
    "Health Care": "XLV", "Industrials": "XLI", "Materials": "XLB",
    "Real Estate": "XLRE", "Technology": "XLK", "Utilities": "XLU",
}


def number(value: Any) -> float | None:
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def compare(value, operation, limit):
    if value is None:
        return False
    return value < limit if operation == "lt" else value > limit


def normalize_ticker(value):
    if value is None or pd.isna(value):
        return None
    value = str(value).strip().upper().replace(".", "-")
    return None if not value or value in {"TICKER", "SYMBOL", "NAN", "NONE"} else value


def create_starter(path):
    wb = Workbook()
    ws = wb.active
    ws.title = TICKERS_SHEET
    ws.append(["Stock", "Dividend", "ETF", "S&P 100"])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 18
    wb.save(path)


def read_tickers(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if TICKERS_SHEET not in wb.sheetnames:
        raise KeyError('The workbook needs a sheet named "Tickers".')
    ws = wb[TICKERS_SHEET]
    result = {}
    for name, col in {"stocks": 1, "dividends": 2, "etfs": 3, "sp100": 4}.items():
        values = [normalize_ticker(ws.cell(row, col).value) for row in range(2, ws.max_row + 1)]
        result[name] = list(dict.fromkeys(v for v in values if v))
    wb.close()
    return result


def download_history(tickers):
    if not tickers:
        return {}
    data = yf.download(tickers, period="1y", interval="1d", auto_adjust=False,
                       actions=False, group_by="ticker", threads=True, progress=False)
    if len(tickers) == 1:
        frame = data.dropna(how="all")
        return {tickers[0]: frame} if not frame.empty else {}
    result = {}
    for ticker in tickers:
        try:
            frame = data[ticker].copy().dropna(how="all")
            if not frame.empty:
                result[ticker] = frame
        except (KeyError, TypeError):
            pass
    return result


def price_metrics(history):
    blank = {
        "Price": None, "Previous Close": None, "Daily Change %": None,
        "1 Month Change %": None, "SMA 50": None, "SMA 200": None,
        "Above SMA 50": False, "Above SMA 200": False,
        "SMA 50 > SMA 200": False, "Bullish Alignment": False,
        "52 Week High": None, "% From 52W High": None,
        "RSI": None, "MACD": None, "MACD Signal": None,
        "MACD Histogram": None, "RSI Test": False, "MACD Test": False,
    }
    if history is None or history.empty:
        return blank
    column = "Adj Close" if "Adj Close" in history.columns else "Close"
    close = pd.to_numeric(history[column], errors="coerce").dropna()
    if close.empty:
        return blank

    price = number(close.iloc[-1])
    previous = number(close.iloc[-2]) if len(close) >= 2 else None
    month_base = number(close.iloc[-(MONTH_LOOKBACK + 1)]) if len(close) > MONTH_LOOKBACK else None
    sma50 = number(close.rolling(50).mean().iloc[-1]) if len(close) >= 50 else None
    sma200 = number(close.rolling(200).mean().iloc[-1]) if len(close) >= 200 else None
    high = number(close.max())
    above50 = price is not None and sma50 is not None and price > sma50
    above200 = price is not None and sma200 is not None and price > sma200
    cross = sma50 is not None and sma200 is not None and sma50 > sma200

    rsi = None
    if len(close) >= 15:
        delta = close.diff()
        gain = delta.clip(lower=0)
        loss = -delta.clip(upper=0)
        avg_gain = gain.ewm(alpha=1 / 14, min_periods=14, adjust=False).mean()
        avg_loss = loss.ewm(alpha=1 / 14, min_periods=14, adjust=False).mean()
        rs = avg_gain / avg_loss.replace(0, float("nan"))
        latest = (100 - 100 / (1 + rs)).iloc[-1]
        if pd.notna(latest):
            rsi = number(latest)
        if rsi is None:
            g, loss_value = number(avg_gain.iloc[-1]), number(avg_loss.iloc[-1])
            if g is not None and g > 0 and loss_value == 0:
                rsi = 100.0
            elif g == 0 and loss_value is not None and loss_value > 0:
                rsi = 0.0

    macd = signal = histogram = None
    if len(close) >= 34:
        ema12 = close.ewm(span=12, adjust=False, min_periods=12).mean()
        ema26 = close.ewm(span=26, adjust=False, min_periods=26).mean()
        macd_series = ema12 - ema26
        signal_series = macd_series.ewm(span=9, adjust=False, min_periods=9).mean()
        macd, signal = number(macd_series.iloc[-1]), number(signal_series.iloc[-1])
        histogram = number((macd_series - signal_series).iloc[-1])

    rsi_test = rsi is not None and MIN_RSI <= rsi <= MAX_RSI
    macd_test = macd is not None and signal is not None and macd > signal
    return {
        "Price": price, "Previous Close": previous,
        "Daily Change %": price / previous - 1 if price is not None and previous not in (None, 0) else None,
        "1 Month Change %": price / month_base - 1 if price is not None and month_base not in (None, 0) else None,
        "SMA 50": sma50, "SMA 200": sma200, "Above SMA 50": above50,
        "Above SMA 200": above200, "SMA 50 > SMA 200": cross,
        "Bullish Alignment": above50 and above200 and cross,
        "52 Week High": high,
        "% From 52W High": price / high - 1 if price is not None and high not in (None, 0) else None,
        "RSI": rsi, "MACD": macd, "MACD Signal": signal,
        "MACD Histogram": histogram, "RSI Test": rsi_test, "MACD Test": macd_test,
    }


def momentum_scores(metrics):
    rsi, macd, signal = number(metrics.get("RSI")), number(metrics.get("MACD")), number(metrics.get("MACD Signal"))
    rsi_score = 2 if rsi is not None and 50 <= rsi <= MAX_RSI else 1 if rsi is not None and MIN_RSI <= rsi < 50 else 0
    macd_score = int(macd is not None and signal is not None and macd > signal)
    return rsi_score, macd_score, rsi_score + macd_score


def get_info(ticker, cache):
    if ticker in cache:
        return cache[ticker]
    try:
        cache[ticker] = yf.Ticker(ticker).get_info() or {}
    except Exception as exc:
        print(f"  Warning: information unavailable for {ticker}: {exc}")
        cache[ticker] = {}
    time.sleep(PAUSE)
    return cache[ticker]


def normalized_dividend_values(info):
    dividend_yield, payout = number(info.get("dividendYield")), number(info.get("payoutRatio"))
    if dividend_yield is not None and dividend_yield > .25:
        dividend_yield /= 100
    if payout is not None and payout > 1:
        payout /= 100
    return dividend_yield, payout


def momentum_failures(failed, p):
    if p["RSI"] is None:
        failed.append("RSI unavailable")
    elif p["RSI"] < MIN_RSI:
        failed.append(f"RSI below {MIN_RSI:g}")
    elif p["RSI"] > MAX_RSI:
        failed.append(f"RSI above {MAX_RSI:g}")
    if p["MACD"] is None or p["MACD Signal"] is None:
        failed.append("MACD unavailable")
    elif not p["MACD Test"]:
        failed.append("MACD not above signal")


def stock_rows(tickers, histories, cache):
    rows = []
    for n, ticker in enumerate(tickers, 1):
        print(f"Stock {n}/{len(tickers)}: {ticker}")
        info, p = get_info(ticker, cache), price_metrics(histories.get(ticker))
        rsi_score, macd_score, momentum_score = momentum_scores(p)
        target = number(info.get("targetMeanPrice"))
        current = p["Price"] or number(info.get("currentPrice"))
        upside = target / current - 1 if target is not None and current not in (None, 0) else None
        values = {
            "PEG": number(info.get("pegRatio") or info.get("trailingPegRatio")),
            "Forward PE": number(info.get("forwardPE")),
            "Return On Equity": number(info.get("returnOnEquity")),
            "Profit Margin": number(info.get("profitMargins")),
            "Revenue Growth": number(info.get("revenueGrowth")),
            "Earnings Growth": number(info.get("earningsGrowth")),
            "Debt To Equity": number(info.get("debtToEquity")), "Upside %": upside,
        }
        specs = [
            ("PEG", "lt", MAX_PEG, "PEG unavailable", "PEG too high"),
            ("Forward PE", "lt", MAX_FORWARD_PE, "Forward PE unavailable", "Forward PE too high"),
            ("Return On Equity", "gt", MIN_ROE, "ROE unavailable", "ROE too low"),
            ("Profit Margin", "gt", MIN_MARGIN, "Profit margin unavailable", "Profit margin too low"),
            ("Revenue Growth", "gt", 0, "Revenue growth unavailable", "Revenue growth not positive"),
            ("Earnings Growth", "gt", 0, "Earnings growth unavailable", "Earnings growth not positive"),
            ("Debt To Equity", "lt", MAX_DEBT_EQUITY, "Debt to equity unavailable", "Debt to equity too high"),
            ("Upside %", "gt", MIN_UPSIDE, "Target upside unavailable", "Target upside not positive"),
        ]
        tests, failed = {}, []
        for key, op, limit, unavailable, message in specs:
            passed = compare(values[key], op, limit)
            tests[key] = passed
            if values[key] is None: failed.append(unavailable)
            elif not passed: failed.append(message)
        above50, above200, golden = p["Above SMA 50"], p["Above SMA 200"], p["SMA 50 > SMA 200"]
        if not above50: failed.append("Below SMA 50")
        if not above200: failed.append("Below SMA 200")
        if not golden: failed.append("SMA 50 not above SMA 200")
        momentum_failures(failed, p)
        base_tests = list(tests.values()) + [above50, above200, golden]
        base_score, max_score = sum(base_tests), len(base_tests) + 3
        total_score = base_score + momentum_score
        rows.append({
            "Ticker": ticker, "Name": info.get("shortName") or info.get("longName"),
            "Sector": info.get("sector"), "Industry": info.get("industry"), **p,
            "Target Price": target, **values, "Free Cash Flow": number(info.get("freeCashflow")),
            **{f"{key} Test": passed for key, passed in tests.items()},
            "Above SMA 50 Test": above50, "Above SMA 200 Test": above200,
            "Golden Cross Test": golden, "Trend Test": above50 and above200 and golden,
            "RSI Range Test": p["RSI Test"], "MACD Above Signal Test": p["MACD Test"],
            "Momentum Test": p["RSI Test"] and p["MACD Test"],
            "Base Score": base_score, "RSI Score": rsi_score, "MACD Score": macd_score,
            "Momentum Score": momentum_score, "Failed Tests": ", ".join(failed) if failed else None,
            "Score": total_score, "Max Score": max_score, "Score %": total_score / max_score,
            "Passes Screen": all(tests.values()) and above200 and momentum_score >= 2,
            "Error": None if info else "Data unavailable",
        })
    return rows


def dividend_rows(tickers, histories, cache):
    rows = []
    for n, ticker in enumerate(tickers, 1):
        print(f"Dividend {n}/{len(tickers)}: {ticker}")
        info, p = get_info(ticker, cache), price_metrics(histories.get(ticker))
        rsi_score, macd_score, momentum_score = momentum_scores(p)
        dividend_yield, payout = normalized_dividend_values(info)
        if dividend_yield is not None and dividend_yield > 0.25:
            dividend_yield = dividend_yield / 100
        cash_flow, debt = number(info.get("freeCashflow")), number(info.get("debtToEquity"))
        average_yield = number(info.get("fiveYearAvgDividendYield"))
        test_values = [
            (dividend_yield, compare(dividend_yield, "gt", MIN_DIVIDEND_YIELD), "Dividend yield unavailable", "Dividend yield below minimum"),
            (payout, compare(payout, "lt", MAX_PAYOUT), "Payout ratio unavailable", "Payout ratio too high"),
            (cash_flow, cash_flow is not None and cash_flow > 0, "Free cash flow unavailable", "Free cash flow not positive"),
            (debt, compare(debt, "lt", MAX_DEBT_EQUITY), "Debt to equity unavailable", "Debt to equity too high"),
        ]
        core_tests = [x[1] for x in test_values]
        failed = [x[2] if x[0] is None else x[3] for x in test_values if x[0] is None or not x[1]]
        above50, above200, golden = p["Above SMA 50"], p["Above SMA 200"], p["SMA 50 > SMA 200"]
        if not above50: failed.append("Below SMA 50")
        if not above200: failed.append("Below SMA 200")
        if not golden: failed.append("SMA 50 not above SMA 200")
        momentum_failures(failed, p)
        base_tests = core_tests + [above50, above200, golden]
        base_score, max_score = sum(base_tests), len(base_tests) + 3
        total_score = base_score + momentum_score
        rows.append({
            "Ticker": ticker, "Name": info.get("shortName") or info.get("longName"), "Sector": info.get("sector"), **p,
            "Dividend Yield": dividend_yield, "5Y Avg Dividend Yield": average_yield / 100 if average_yield is not None else None,
            "Payout Ratio": payout, "Free Cash Flow": cash_flow, "Debt To Equity": debt,
            "Dividend Yield Test": core_tests[0], "Payout Ratio Test": core_tests[1],
            "Free Cash Flow Test": core_tests[2], "Debt To Equity Test": core_tests[3],
            "Above SMA 50 Test": above50, "Above SMA 200 Test": above200,
            "Golden Cross Test": golden, "Trend Test": above50 and above200 and golden,
            "RSI Range Test": p["RSI Test"], "MACD Above Signal Test": p["MACD Test"],
            "Momentum Test": p["RSI Test"] and p["MACD Test"],
            "Base Score": base_score, "RSI Score": rsi_score, "MACD Score": macd_score,
            "Momentum Score": momentum_score, "Failed Tests": ", ".join(failed) if failed else None,
            "Score": total_score, "Max Score": max_score, "Score %": total_score / max_score,
            "Passes Screen": all(core_tests) and above200 and momentum_score >= 2,
            "Error": None if info else "Data unavailable",
        })
    return rows


def get_etf_metrics(ticker, info, cache):
    if ticker in cache: return cache[ticker]
    expense = number(info.get("annualReportExpenseRatio") or info.get("netExpenseRatio"))
    assets = number(info.get("totalAssets"))
    cache[ticker] = {"Expense Ratio": expense, "Total Assets": assets}
    return cache[ticker]


def etf_rows(tickers, histories, cache, fund_cache):
    rows = []
    for n, ticker in enumerate(tickers, 1):
        print(f"ETF {n}/{len(tickers)}: {ticker}")
        info, p = get_info(ticker, cache), price_metrics(histories.get(ticker))
        fund = get_etf_metrics(ticker, info, fund_cache)
        expense, assets = fund["Expense Ratio"], fund["Total Assets"]
        tests = [compare(expense, "lt", MAX_EXPENSE_RATIO), compare(assets, "gt", MIN_AUM), p["Above SMA 50"], p["Above SMA 200"], p["SMA 50 > SMA 200"]]
        failed = []
        if expense is None: failed.append("Expense ratio unavailable")
        elif not tests[0]: failed.append("Expense ratio too high")
        if assets is None: failed.append("AUM unavailable")
        elif not tests[1]: failed.append("AUM below minimum")
        if not tests[2]: failed.append("Below SMA 50")
        if not tests[3]: failed.append("Below SMA 200")
        if not tests[4]: failed.append("SMA 50 not above SMA 200")
        rows.append({"Ticker": ticker, "Name": info.get("shortName") or info.get("longName"), "Category": info.get("category"), **p,
                     "Expense Ratio": expense, "Total Assets": assets, "Expense Test": tests[0], "AUM Test": tests[1],
                     "Trend Test": all(tests[2:]), "Failed Tests": ", ".join(failed) if failed else None,
                     "Yield": number(info.get("yield")), "3Y Average Return": number(info.get("threeYearAverageReturn")),
                     "5Y Average Return": number(info.get("fiveYearAverageReturn")), "Score": sum(tests), "Max Score": len(tests),
                     "Passes Screen": all(tests), "Error": None if info else "Data unavailable"})
    return rows


def heatmap_rows(tickers, histories, cache, dividend=False):
    rows = []
    for ticker in tickers:
        info, p = get_info(ticker, cache), price_metrics(histories.get(ticker))
        row = {"Ticker": ticker, "Name": info.get("shortName") or info.get("longName"),
               "Group": info.get("sector") or info.get("category") or "Unclassified",
               "Market Cap / Assets": number(info.get("marketCap") or info.get("totalAssets")) or 1,
               "Price": p["Price"], "Previous Close": p["Previous Close"], "Daily Change %": p["Daily Change %"], "1 Month Change %": p["1 Month Change %"],
               "SMA 50": p["SMA 50"], "SMA 200": p["SMA 200"], "Bullish Alignment": p["Bullish Alignment"]}
        if dividend:
            row["Dividend Yield"], row["Payout Ratio"] = normalized_dividend_values(info)
        rows.append(row)
    return rows


def sector_heatmap_rows(histories, cache):
    rows = []
    for sector, ticker in SECTOR_ETFS.items():
        info, p = get_info(ticker, cache), price_metrics(histories.get(ticker))
        rows.append({"Sector": sector, "Ticker": ticker, "Name": info.get("shortName") or sector,
                     "Weight": 1, "Price": p["Price"], "Previous Close": p["Previous Close"],
                     "Daily Change %": p["Daily Change %"], "1 Month Change %": p["1 Month Change %"],
                     "SMA 50": p["SMA 50"], "SMA 200": p["SMA 200"], "Bullish Alignment": p["Bullish Alignment"]})
    return rows


def sort_frame(rows, columns, ascending):
    frame = pd.DataFrame(rows)
    return frame.sort_values(columns, ascending=ascending, na_position="last").reset_index(drop=True) if not frame.empty else frame


def create_treemap(frame, title, sector=False):
    if frame.empty:
        return None
    chart = frame.copy()
    value_col = "Weight" if sector else "Market Cap / Assets"
    group_col = "Sector" if sector else "Group"
    chart[value_col] = pd.to_numeric(chart[value_col], errors="coerce").fillna(1).clip(lower=1)
    chart["1 Month Change %"] = pd.to_numeric(chart["1 Month Change %"], errors="coerce")
    chart = chart.dropna(subset=["1 Month Change %"])
    if chart.empty:
        return None

    # Descriptive hover details. Plotly includes only fields available for each map.
    hover_data = {}
    hover_formats = {
        "Name": True,
        "Price": ":$,.2f",
        "Previous Close": ":$,.2f",
        "Daily Change %": ":+.2%",
        "1 Month Change %": ":+.2%",
        "SMA 50": ":$,.2f",
        "SMA 200": ":$,.2f",
        "Bullish Alignment": True,
        "Dividend Yield": ":.2%",
        "Payout Ratio": ":.1%",
    }
    for field, display in hover_formats.items():
        if field in chart.columns:
            hover_data[field] = display

    fig = px.treemap(
        chart,
        path=[px.Constant("Market"), group_col, "Ticker"],
        values=value_col,
        color="1 Month Change %",
        color_continuous_midpoint=0,
        color_continuous_scale=[
            [0, "#8B0000"], [.25, "#D73027"], [.45, "#F4A6A6"],
            [.5, "#333333"], [.55, "#A7DCA7"], [.75, "#1A9850"], [1, "#006400"],
        ],
        hover_data=hover_data,
        title=title,
    )

    # Show the one-month increase/decrease directly on each heatmap tile.
    labels = [
        f"{number(value):+.2%}" if number(value) is not None else ""
        for value in fig.data[0].marker.colors
    ]
    fig.update_traces(
        text=labels,
        texttemplate="<b>%{label}</b><br>%{text}",
        marker=dict(line=dict(width=1, color="#202020")),
        root_color="#101820",
    )
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="#101820",
        plot_bgcolor="#101820",
        margin=dict(t=55, l=8, r=8, b=8),
        coloraxis_colorbar=dict(title="1 Month Change", tickformat="+.1%"),
    )
    return fig

def export_assets(folder, frames):
    figures = [("S&P 100 Map", create_treemap(frames["Stock Heatmap Data"], "S&P 100 One-Month Heatmap"), "sp100_heatmap.png"),
               ("Dividend Map", create_treemap(frames["Dividend Heatmap Data"], "Dividend Aristocrat One-Month Heatmap"), "dividend_heatmap.png"),
               ("ETF Map", create_treemap(frames["ETF Heatmap Data"], "ETF One-Month Heatmap"), "etf_heatmap.png"),
               ("Sector Map", create_treemap(frames["Sector Heatmap Data"], "U.S. Sector ETF One-Month Heatmap", True), "sector_heatmap.png")]
    images, html_parts = {}, []
    for index, (name, fig, filename) in enumerate(figures):
        if fig:
            html_parts.append(fig.to_html(full_html=False, include_plotlyjs="inline" if index == 0 else False))
            try:
                image_path = folder / filename
                fig.write_image(image_path, width=1800, height=1000)
                images[name] = image_path
            except Exception as exc:
                print(f"Warning: PNG export failed for {name}: {exc}")
    html = "<!doctype html><html><head><meta charset='utf-8'><title>Investment Dashboard</title></head><body style='background:#101820;color:white'>" + "".join(html_parts) + "</body></html>"
    (folder / HTML_NAME).write_text(html, encoding="utf-8")
    return images


def write_dataframe(workbook, name, frame):
    if name in workbook.sheetnames: del workbook[name]
    ws = workbook.create_sheet(name)
    ws.freeze_panes, ws.sheet_view.showGridLines = "A2", False
    if frame.empty:
        ws["A1"] = "No results"; return
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None): ws.append(list(row))
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = ws.dimensions
    headers = {cell.value: cell.column for cell in ws[1]}
    percentages = {"Daily Change %", "1 Month Change %", "% From 52W High", "Upside %", "Return On Equity", "Profit Margin", "Revenue Growth", "Earnings Growth", "Dividend Yield", "5Y Avg Dividend Yield", "Payout Ratio", "Expense Ratio", "Yield", "3Y Average Return", "5Y Average Return", "Score %"}
    currencies = {"Price", "Previous Close", "SMA 50", "SMA 200", "52 Week High", "Target Price"}
    for header, col in headers.items():
        fmt = "0.0%" if header in percentages else "$#,##0.00" if header in currencies else None
        if fmt:
            for cell in list(ws.columns)[col - 1][1:]: cell.number_format = fmt
    for header in ("Daily Change %", "1 Month Change %"):
        if header in headers:
            letter = get_column_letter(headers[header])
            ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    for cells in ws.columns:
        ws.column_dimensions[get_column_letter(cells[0].column)].width = min(max(max(len(str(c.value or "")) for c in cells[:200]) + 2, 11), 32)


def insert_map(workbook, name, image_path, html_path):
    if name in workbook.sheetnames: del workbook[name]
    ws = workbook.create_sheet(name); ws["A1"] = name; ws["A2"] = "Open interactive dashboard"; ws["A2"].hyperlink = html_path.resolve().as_uri(); ws["A2"].style = "Hyperlink"
    if image_path and image_path.exists():
        image = XLImage(str(image_path)); image.width = 1350; image.height = 750; ws.add_image(image, "A4")


def save_workbook(path, frames, images):
    wb = load_workbook(path)
    for name, frame in frames.items(): write_dataframe(wb, name, frame)
    for name in ["S&P 100 Map", "Dividend Map", "ETF Map", "Sector Map"]: insert_map(wb, name, images.get(name), path.parent / HTML_NAME)
    if "Dashboard" in wb.sheetnames: del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0); ws.append(["Investment Screening Dashboard"]); ws.append(["Last Refresh", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")]); ws.append([]); ws.append(["Section", "Tickers", "Passing", "Strong Uptrends"])
    for label, key in [("Stocks", "Stock Screener"), ("Dividends", "Dividend Screener"), ("ETFs", "ETF Screener")]:
        frame = frames[key]; passing = int(frame["Passes Screen"].fillna(False).sum()) if not frame.empty else 0; bullish = int(frame["Bullish Alignment"].fillna(False).sum()) if not frame.empty else 0; ws.append([label, len(frame), passing, bullish])
    wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.save(path)


def main():
    folder = Path(__file__).resolve().parent
    workbook_path = folder / WORKBOOK_NAME
    if not workbook_path.exists():
        create_starter(workbook_path); print(f"Created {workbook_path}. Add tickers and run again."); return
    lists = read_tickers(workbook_path)
    all_tickers = list(dict.fromkeys(lists["stocks"] + lists["dividends"] + lists["etfs"] + lists["sp100"] + list(SECTOR_ETFS.values())))
    if not all_tickers: print("No tickers found."); return
    histories, cache, fund_cache = download_history(all_tickers), {}, {}
    frames = {
        "Stock Screener": sort_frame(stock_rows(lists["stocks"], histories, cache), ["Passes Screen", "Score", "Ticker"], [False, False, True]),
        "Dividend Screener": sort_frame(dividend_rows(lists["dividends"], histories, cache), ["Passes Screen", "Score", "Dividend Yield"], [False, False, False]),
        "ETF Screener": sort_frame(etf_rows(lists["etfs"], histories, cache, fund_cache), ["Passes Screen", "Score", "Ticker"], [False, False, True]),
        "Stock Heatmap Data": sort_frame(heatmap_rows(lists["sp100"], histories, cache), ["Group", "1 Month Change %"], [True, False]),
        "Dividend Heatmap Data": sort_frame(heatmap_rows(lists["dividends"], histories, cache, True), ["Group", "1 Month Change %"], [True, False]),
        "ETF Heatmap Data": sort_frame(heatmap_rows(lists["etfs"], histories, cache), ["Group", "1 Month Change %"], [True, False]),
        "Sector Heatmap Data": sort_frame(sector_heatmap_rows(histories, cache), ["Daily Change %", "Sector"], [False, True]),
    }
    images = export_assets(folder, frames); save_workbook(workbook_path, frames, images)
    print(f"Updated workbook: {workbook_path}"); print(f"Interactive dashboard: {folder / HTML_NAME}")


if __name__ == "__main__":
    main()
