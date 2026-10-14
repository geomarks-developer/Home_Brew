import sys
import datetime
import pandas as pd
import yfinance as yf
import os
import numpy as np
from openpyxl import load_workbook
from pathlib import Path

# --- Choose your tickers --------------------------------------------------
# Note: yfinance wants dots replaced with dashes, e.g. "BRK.B" -> "BRK-B"
#TICKERS = ["AAPL","ABBV","ABT","ACN","ADBE","AEP","AIG","ALL","AMGN","AMT","AMZN"
#,"AXP","BA","BAC","BK","BKNG","BLK","BMY","BRK-B","C","CAT","CHTR","CL","CMCSA"
#,"COF","COP","COST","CRM","CSCO","CVS","CVX","D","DE","DHR","DIS","DUK","EMR","EXC"
#,"FDX","GD","GE","GILD","GM","GOOG","GOOGL","GS","HD","HON","IBM","INTC","JNJ","JPM"
#,"KHC","KMI","KO","LIN","LLY","LMT","LOW","MA","MCD","MDLZ","MET","META","MMM","MO"
#,"MRK","MS","MSFT","NEE","NFLX","NKE","NVDA","ORCL","PEP","PFE","PG","PM","PYPL","QCOM"
#,"RTX","SBUX","SCHW","SO","SPGI","T","TGT","TMO","TMUS","TSLA","TXN","UNH","UNP","UPS"
#,"USB","V","VZ","WBA","WFC","WMT","XOM"
#]

base_dir = Path(__file__).resolve().parent
excel_path = base_dir / "StockAnalysis.xlsx"

wb = load_workbook(excel_path)
ws = wb["Tickers"]

TICKERS = [
    str(cell.value).strip().upper()
    for cell in ws["A"][1:]      # skip header
    if cell.value
]

#print(TICKERS)

# --- Choose your columns ---------------------------------------------------
# Keys come from yfinance's Ticker.info dictionary. Add/remove freely.
# Give each one a friendly display name on the right.
FIELDS = {
    "shortName": "Name",
    "sector": "Sector",
    "previousClose": "Previous Close",
    "currentPrice": "Last Price",
    "targetMeanPrice": "Target Price",
    "marketCap": "Market Cap",
    #"trailingPE": "P/E (TTM)",
    "pegRatio": "PEG",
    #"dividendYield": "Dividend Yield",
    #"beta": "Beta (5Y)",
    "forwardPE": "Forward PE",
    #"priceToBook": "Price To Book",
    "revenueGrowth": "Revenue Growth",
    "earningsGrowth": "Earnings Growth",
    "profitMargins": "Profit Margin",
    "returnOnEquity": "Return On Equity",
    "debtToEquity": "Debt To Equity",
    "freeCashflow": "Free Cash Flow",
    #"operatingCashflow": "Operating Cash Flow",
    "fiftyDayAverage": "50 Day MA",
    "twoHundredDayAverage": "200 Day MA"
}

def fetch_selected_fields(symbol: str, fields: dict) -> dict:
    """Pull only the named fields from yfinance's info dict for one ticker."""
    info = yf.Ticker(symbol).info
    row = {"Ticker": symbol}
    for key, display_name in fields.items():
        row[display_name] = info.get(key)  # None if the field isn't available for this ticker
    return row

def safe_num(value, default=0):
    """Convert None/NaN to a usable number."""
    if value is None or pd.isna(value):
        return default
    return value

def calculate_metrics(df):
    """Add Upside %, Value Score, and Screen Result."""

    # Convert columns to numeric
    numeric_cols = [
        "Previous Close",
        "Last Price",
        "Target Price",
        "PEG",
        "Forward PE",
        "Revenue Growth",
        "Earnings Growth",
        "Profit Margin",
        "Return On Equity",
        "Debt To Equity",
        "Free Cash Flow"
    ]

    for col in numeric_cols:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    # Analyst upside
    df["Upside %"] = (
        (df["Target Price"] - df["Last Price"])
        / df["Last Price"]
    )

    # Replace invalid values
    df["PEG"] = df["PEG"].replace(0, np.nan)
    df["Forward PE"] = df["Forward PE"].replace(0, np.nan)

    # Value Score Formula
    df["Value Score"] = (
        0.30 * df["Upside %"].fillna(0)
        + 0.20 * (1 / df["PEG"].fillna(999))
        + 0.15 * (1 / df["Forward PE"].fillna(999))
        + 0.10 * df["Return On Equity"].fillna(0)
        + 0.10 * df["Profit Margin"].fillna(0)
        + 0.10 * df["Earnings Growth"].fillna(0)
        + 0.05 * df["Revenue Growth"].fillna(0)
    )

    # Screening filters
    df["Passes Screen"] = (
        (df["PEG"] < 2.5)
        & (df["Forward PE"] < 25)
        & (df["Return On Equity"] > 0.10)
        & (df["Profit Margin"] > 0.08)
        & (df["Revenue Growth"] > 0.05)
        & (df["Earnings Growth"] > 0.05)
        & (df["Debt To Equity"] < 100)
        & (df["Free Cash Flow"] > 0)
        & (df["Upside %"] > 0.20)
    )

    # Rank by score
    df = df.sort_values(
        by="Value Score",
        ascending=False
    )

    df["Rating"] = np.select(
        [
            df["Value Score"] >= 1.0,
            df["Value Score"] >= 0.5,
            df["Value Score"] >= 0.25
        ],
        [
            "Strong Buy",
            "Attractive",
            "Watch List"
        ],
        default="Pass"
    )

    df = df.sort_values(
    by="Value Score",
    ascending=False
    )

    return df

def main(tickers):
    rows = []
    for symbol in tickers:
        print(f"Fetching {symbol}...")
        try:
            rows.append(fetch_selected_fields(symbol, FIELDS))
        except Exception as exc:
            print(f"  ! Could not fetch {symbol}: {exc}")
            rows.append({"Ticker": symbol, **{name: None for name in FIELDS.values()}})

    df = pd.DataFrame(rows)

    # Calculate scores and rankings
    df = calculate_metrics(df)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    screened_df = df[df["Passes Screen"]]

    if os.path.exists(excel_path):

        with pd.ExcelWriter(
            excel_path,
            engine="openpyxl",
            mode="a",
            if_sheet_exists="replace"
        ) as writer:

            df.to_excel(writer, sheet_name="Data", index=False)
            screened_df.to_excel(writer, sheet_name="Screened", index=False)

    else:

        with pd.ExcelWriter(
            excel_path,
            engine="openpyxl",
            mode="w"
        ) as writer:

            df.to_excel(writer, sheet_name="Data", index=False)
            screened_df.to_excel(writer, sheet_name="Screened", index=False)

    print(f"Saved results to {excel_path}")

if __name__ == "__main__":
    tickers_arg = sys.argv[1:] if len(sys.argv) > 1 else TICKERS
    main(tickers_arg)